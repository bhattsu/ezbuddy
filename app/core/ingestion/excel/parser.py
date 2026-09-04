"""Parse legal configuration Excel workbooks into typed models (no DB I/O)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, TypeVar

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from pydantic import BaseModel, ValidationError

from app.core.ingestion.excel.constants import (
    DOCUMENT_COLUMNS,
    KNOWLEDGE_COLUMNS,
    LOCAL_OVERRIDE_COLUMNS,
    LOOKUP_COLUMNS,
    PACKAGE_COLUMNS,
    QUESTION_COLUMNS,
    SHEET_DOCUMENTS,
    SHEET_KNOWLEDGE,
    SHEET_LOCAL_OVERRIDES,
    SHEET_LOOKUPS,
    SHEET_PACKAGE,
    SHEET_QUESTIONS,
    SHEET_WORKFLOW_MAPPING,
    SHEET_WORKFLOWS,
    WORKFLOW_COLUMNS,
    WORKFLOW_MAPPING_COLUMNS,
)
from app.core.ingestion.excel.exceptions import EmptySheetError, PackageSheetError
from app.core.ingestion.excel.mappers import dedupe_workflow_mappings
from app.core.ingestion.excel.schemas import (
    DocumentRow,
    ExcelPackage,
    KnowledgeRow,
    LocalOverrideRow,
    LookupValueRow,
    PackageRow,
    QuestionRow,
    WorkflowMappingRow,
    WorkflowRow,
)
from app.core.ingestion.excel.validators import (
    clean_cell,
    ignored_excel_columns,
    parse_int,
    parse_yes_no,
    row_is_empty,
    validate_sheet_headers,
    validate_workbook_sheets,
)

ModelT = TypeVar("ModelT", bound=BaseModel)


class LegalPackageExcelParser:
    """Generic parser for US Legal Pro configuration Excel packages."""

    def parse(self, file_path: str | Path) -> ExcelPackage:
        path = Path(file_path)
        workbook = self._load_workbook(path)
        warnings: list[str] = []

        try:
            validate_workbook_sheets(workbook)

            package = self._parse_package(workbook)
            workflows = self._parse_workflows(workbook, warnings)
            questions = self._parse_questions(workbook, warnings)
            lookups = self._parse_lookups(workbook, warnings)
            mappings = dedupe_workflow_mappings(
                self._parse_workflow_mappings(workbook, warnings),
                warnings,
            )
            documents = self._parse_documents(workbook, warnings)
            knowledge = self._parse_knowledge(workbook, warnings)
            overrides = self._parse_local_overrides(workbook, warnings)
            ignored_columns = self._collect_ignored_columns(workbook)

            return ExcelPackage(
                package=package,
                workflows=workflows,
                questions=questions,
                lookup_values=lookups,
                workflow_mappings=mappings,
                documents=documents,
                knowledge=knowledge,
                local_overrides=overrides,
                source_path=str(path.resolve()),
                parse_warnings=warnings,
                ignored_columns=ignored_columns,
            )
        finally:
            workbook.close()

    def _load_workbook(self, path: Path) -> Workbook:
        if not path.is_file():
            raise FileNotFoundError(path)
        return load_workbook(path, read_only=True, data_only=True)

    @staticmethod
    def _collect_ignored_columns(workbook: Workbook) -> dict[str, list[str]]:
        from app.core.ingestion.excel.constants import SHEET_COLUMN_MAP

        out: dict[str, list[str]] = {}
        for sheet_name in SHEET_COLUMN_MAP:
            if sheet_name not in workbook.sheetnames:
                continue
            ws = workbook[sheet_name]
            header = next(ws.iter_rows(values_only=True), None)
            if header is None:
                continue
            extras = ignored_excel_columns(sheet_name, header)
            if extras:
                out[sheet_name] = extras
        return out

    def _parse_package(self, workbook: Workbook) -> PackageRow:
        ws = workbook[SHEET_PACKAGE]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise EmptySheetError(SHEET_PACKAGE, "no header row")

        header_idx = validate_sheet_headers(SHEET_PACKAGE, header)
        data_row = next(rows, None)
        while data_row is not None and row_is_empty(data_row):
            data_row = next(rows, None)
        if data_row is None or row_is_empty(data_row):
            raise PackageSheetError("Package sheet must contain exactly one data row")

        raw = self._row_to_dict(data_row, header_idx, PACKAGE_COLUMNS)
        try:
            return PackageRow.model_validate(raw)
        except ValidationError as exc:
            raise PackageSheetError(f"invalid Package row: {exc}") from exc

    def _parse_workflows(self, workbook: Workbook, warnings: list[str]) -> list[WorkflowRow]:
        return self._parse_sheet(
            workbook,
            SHEET_WORKFLOWS,
            WORKFLOW_COLUMNS,
            WorkflowRow,
            warnings,
            transform=self._transform_workflow,
        )

    def _parse_questions(self, workbook: Workbook, warnings: list[str]) -> list[QuestionRow]:
        return self._parse_sheet(
            workbook,
            SHEET_QUESTIONS,
            QUESTION_COLUMNS,
            QuestionRow,
            warnings,
            transform=self._transform_question,
        )

    def _parse_lookups(self, workbook: Workbook, warnings: list[str]) -> list[LookupValueRow]:
        return self._parse_sheet(
            workbook,
            SHEET_LOOKUPS,
            LOOKUP_COLUMNS,
            LookupValueRow,
            warnings,
            transform=self._transform_lookup,
        )

    def _parse_workflow_mappings(
        self, workbook: Workbook, warnings: list[str]
    ) -> list[WorkflowMappingRow]:
        return self._parse_sheet(
            workbook,
            SHEET_WORKFLOW_MAPPING,
            WORKFLOW_MAPPING_COLUMNS,
            WorkflowMappingRow,
            warnings,
            transform=self._transform_mapping,
            required_fields={"workflow_code", "question_code", "display_order"},
        )

    def _parse_documents(self, workbook: Workbook, warnings: list[str]) -> list[DocumentRow]:
        return self._parse_sheet(
            workbook,
            SHEET_DOCUMENTS,
            DOCUMENT_COLUMNS,
            DocumentRow,
            warnings,
            transform=self._transform_document,
        )

    def _parse_knowledge(self, workbook: Workbook, warnings: list[str]) -> list[KnowledgeRow]:
        return self._parse_sheet(
            workbook,
            SHEET_KNOWLEDGE,
            KNOWLEDGE_COLUMNS,
            KnowledgeRow,
            warnings,
            transform=self._transform_knowledge,
        )

    def _parse_local_overrides(
        self, workbook: Workbook, warnings: list[str]
    ) -> list[LocalOverrideRow]:
        return self._parse_sheet(
            workbook,
            SHEET_LOCAL_OVERRIDES,
            LOCAL_OVERRIDE_COLUMNS,
            LocalOverrideRow,
            warnings,
            transform=self._transform_override,
            allow_empty=True,
        )

    def _parse_sheet(
        self,
        workbook: Workbook,
        sheet_name: str,
        column_map: dict[str, str],
        model: type[ModelT],
        warnings: list[str],
        *,
        transform: Callable[[dict[str, Any]], dict[str, Any]] | None = None,
        required_fields: set[str] | None = None,
        allow_empty: bool = False,
    ) -> list[ModelT]:
        ws = workbook[sheet_name]
        row_iter = ws.iter_rows(values_only=True)
        header = next(row_iter, None)
        if header is None:
            if allow_empty:
                return []
            raise EmptySheetError(sheet_name, "no header row")

        header_idx = validate_sheet_headers(sheet_name, header)
        results: list[ModelT] = []
        row_num = 1

        for data_row in row_iter:
            row_num += 1
            if row_is_empty(data_row):
                continue
            raw = self._row_to_dict(data_row, header_idx, column_map)
            if transform:
                raw = transform(raw)
            if required_fields:
                missing = [f for f in required_fields if raw.get(f) in (None, "")]
                if missing:
                    warnings.append(
                        f"{sheet_name} row {row_num}: skipped, missing {missing}"
                    )
                    continue
            try:
                results.append(model.model_validate(raw))
            except ValidationError as exc:
                warnings.append(f"{sheet_name} row {row_num}: skipped, {exc.errors()[0]['msg']}")
        return results

    @staticmethod
    def _row_to_dict(
        row: tuple[Any, ...],
        header_idx: dict[str, int],
        column_map: dict[str, str],
    ) -> dict[str, Any]:
        out: dict[str, Any] = {}
        for excel_header, field_name in column_map.items():
            idx = header_idx.get(excel_header)
            if idx is None or idx >= len(row):
                out[field_name] = None
            else:
                out[field_name] = clean_cell(row[idx])
        return out

    @staticmethod
    def _transform_workflow(raw: dict[str, Any]) -> dict[str, Any]:
        raw["reusable"] = parse_yes_no(raw.get("reusable"))
        return raw

    @staticmethod
    def _transform_question(raw: dict[str, Any]) -> dict[str, Any]:
        raw["reusable"] = parse_yes_no(raw.get("reusable"))
        return raw

    @staticmethod
    def _transform_lookup(raw: dict[str, Any]) -> dict[str, Any]:
        raw["display_order"] = parse_int(raw.get("display_order"))
        raw["active"] = parse_yes_no(raw.get("active"))
        return raw

    @staticmethod
    def _transform_mapping(raw: dict[str, Any]) -> dict[str, Any]:
        raw["display_order"] = parse_int(raw.get("display_order"))
        if raw.get("display_order") is None:
            raw["display_order"] = 0
        raw["page"] = parse_int(raw.get("page"))
        raw["required"] = parse_yes_no(raw.get("required"))
        raw["is_editable"] = parse_yes_no(raw.get("is_editable"))
        return raw

    @staticmethod
    def _transform_document(raw: dict[str, Any]) -> dict[str, Any]:
        raw["required"] = parse_yes_no(raw.get("required"))
        return raw

    @staticmethod
    def _transform_knowledge(raw: dict[str, Any]) -> dict[str, Any]:
        for key in ("effective_date", "last_verified"):
            val = raw.get(key)
            if isinstance(val, datetime):
                raw[key] = val.date()
        return raw

    @staticmethod
    def _transform_override(raw: dict[str, Any]) -> dict[str, Any]:
        raw["override_priority"] = parse_int(raw.get("override_priority"))
        val = raw.get("effective_date")
        if isinstance(val, datetime):
            raw["effective_date"] = val.date()
        return raw
