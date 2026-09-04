"""
DOCX / Word Document Loader

Extracts plain text (and tables) from .docx files for RAG ingestion.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import List

from app.core.pipelines.ingestion.loaders.base import Document, DocumentLoader

logger = logging.getLogger(__name__)


class DocxLoader(DocumentLoader):
    """Loader for Microsoft Word .docx files via python-docx."""

    def __init__(self):
        self.supported_extensions = {".docx"}

    async def load(self, source: str) -> List[Document]:
        file_path = Path(source)
        if not file_path.exists():
            logger.error("DOCX file not found: %s", source)
            return []
        if file_path.suffix.lower() not in self.supported_extensions:
            logger.error("Unsupported Word extension for DocxLoader: %s", file_path.suffix)
            return []

        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError(
                "python-docx is required to ingest .docx files"
            ) from exc

        try:
            docx = DocxDocument(str(file_path))
            parts: List[str] = []
            for paragraph in docx.paragraphs:
                text = (paragraph.text or "").strip()
                if text:
                    parts.append(text)

            for table in docx.tables:
                for row in table.rows:
                    cells = [(cell.text or "").strip() for cell in row.cells]
                    row_text = " | ".join(c for c in cells if c)
                    if row_text:
                        parts.append(row_text)

            content = "\n\n".join(parts).strip()
            if not content:
                logger.warning("DOCX file has no extractable text: %s", source)
                return []

            metadata = self._get_file_metadata(file_path)
            metadata["loader"] = "docx"
            return [
                Document(
                    content=content,
                    metadata=metadata,
                    source=str(file_path),
                )
            ]
        except Exception:
            logger.exception("Failed to load DOCX file: %s", source)
            return []

    async def load_directory(self, directory: str, pattern: str = "*.docx") -> List[Document]:
        directory_path = Path(directory)
        if not directory_path.is_dir():
            logger.error("Directory not found: %s", directory)
            return []

        documents: List[Document] = []
        for file_path in sorted(directory_path.glob(pattern)):
            if file_path.is_file():
                documents.extend(await self.load(str(file_path)))
        return documents
