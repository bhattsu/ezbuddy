import asyncio
from typing import List, Optional
import logging
import pandas as pd
import openpyxl
from concurrent.futures import ThreadPoolExecutor

from app.api.schemas.document import PageData, TableInfo
from app.config.settings import settings
from app.utils.table_utils import create_table_info

logger = logging.getLogger(__name__)

class AsyncExcelExtractor:
    """Async Excel extraction with openpyxl and pandas"""

    def __init__(self):
        self.executor = ThreadPoolExecutor(max_workers=settings.MAX_WORKERS)

    async def extract_sheets(
        self,
        file_path: str,
        document_id: str,
        sheets_to_process: Optional[List[str]] = None,
        extract_text: bool = True,
        extract_tables: bool = True
    ) -> List[PageData]:
        """Extract data from Excel sheets concurrently"""
        loop = asyncio.get_event_loop()
        sheet_names = await loop.run_in_executor(
            self.executor,
            self._get_sheet_names,
            file_path
        )
        if sheets_to_process:
            sheet_names = [s for s in sheet_names if s in sheets_to_process]
        if not sheet_names:
            logger.warning("No sheets to process")
            return []
        tasks = [
            loop.run_in_executor(
                self.executor,
                self._extract_single_sheet,
                file_path,
                document_id,
                sheet_name,
                idx + 1,
                extract_text,
                extract_tables
            )
            for idx, sheet_name in enumerate(sheet_names)
        ]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        pages_data = []
        for result in results:
            if isinstance(result, Exception):
                logger.error(f"Sheet extraction failed: {str(result)}")
                continue
            if result:
                pages_data.append(result)
        return pages_data

    def _get_sheet_names(self, file_path: str) -> List[str]:
        try:
            workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            return sheet_names
        except Exception as e:
            logger.error(f"Failed to get sheet names: {str(e)}")
            raise RuntimeError(f"Failed to read Excel file: {str(e)}")

    def _extract_single_sheet(
        self,
        file_path: str,
        document_id: str,
        sheet_name: str,
        sheet_number: int,
        extract_text: bool,
        extract_tables: bool
    ) -> PageData:
        page_data = PageData(page_number=sheet_number)
        page_data.metadata = {"sheet_name": sheet_name}
        try:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            if extract_text:
                page_data.text_content = df.to_string()
            if extract_tables:
                headers = [str(h) for h in df.columns.tolist()]
                rows = [[str(cell) for cell in row] for row in df.values.tolist()]
                table_info = TableInfo(
                    table_id=f"{document_id}_sheet{sheet_number}",
                    page_number=sheet_number,
                    headers=headers,
                    rows=rows
                )
                page_data.tables.append(table_info)
            logger.info(f"Extracted sheet '{sheet_name}': {len(df)} rows, {len(df.columns)} cols")
        except Exception as e:
            logger.error(f"Failed to extract sheet '{sheet_name}': {str(e)}")
            page_data.metadata['error'] = str(e)
        return page_data
