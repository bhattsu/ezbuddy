import asyncio
from typing import List
import logging
from PIL import Image
import pytesseract 
from docx import Document
import pandas as pd
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor

from app.api.schemas.document import PageData, TableInfo, ImageInfo
from app.config.settings import settings

logger = logging.getLogger(__name__)


class AsyncImageExtractor:
    """Async image OCR extraction using Tesseract (pytesseract)"""

    def __init__(self):
        # If GPU is enabled, avoid multiple workers (GPU memory issues) - kept for compatibility
        max_workers = 1 if getattr(settings, "USE_GPU", False) else settings.MAX_WORKERS
        self.executor = ThreadPoolExecutor(max_workers=max_workers)

        # No heavy model initialization needed for Tesseract

    async def extract_from_image(
        self,
        file_path: str,
        document_id: str
    ) -> List[PageData]:
        """Extract text from image using Tesseract OCR"""
        loop = asyncio.get_running_loop()
        try:
            page_data = await loop.run_in_executor(
                self.executor,
                self._ocr_image,
                file_path,
                document_id
            )
            return [page_data]
        except Exception as e:
            logger.error(f"Image OCR failed: {str(e)}")
            raise RuntimeError(f"Image OCR failed: {str(e)}")

    def _ocr_image(self, file_path: str, document_id: str) -> PageData:
        page_data = PageData(page_number=1)

        try:
            image = Image.open(file_path)

            image_info = ImageInfo(
                image_id=f"{document_id}_img_1",
                file_path=file_path,
                page_number=1,
                format=image.format.lower() if image.format else "unknown",
                width=image.width,
                height=image.height
            )
            page_data.images.append(image_info)

            # Tesseract OCR call
            # Using page segmentation mode 6 (treat image as a single uniform block of text)
            # This approximates EasyOCR's paragraph=True. Adjust --psm if needed.
            text_content = pytesseract.image_to_string(
                image,
                config='--psm 6'  # You can change this based on your document layout
            ).strip()

            page_data.text_content = text_content or "No text detected"

            page_data.metadata = {
                "format": image.format,
                "size": image.size,
                "mode": image.mode,
                "ocr_engine": "tesseract"
            }

            logger.info(f"OCR extracted {len(text_content)} characters")

        except Exception as e:
            logger.error(f"OCR failed: {str(e)}")
            page_data.text_content = f"OCR error: {str(e)}"

        return page_data


class AsyncWordExtractor:
    """Async Word document extraction"""

    def __init__(self):
        self.executor = ThreadPoolExecutor(max_workers=settings.MAX_WORKERS)

    async def extract_from_word(
        self,
        file_path: str,
        document_id: str,
        extract_text: bool = True,
        extract_tables: bool = True
    ) -> List[PageData]:
        """Extract data from Word document"""
        loop = asyncio.get_running_loop()
        try:
            page_data = await loop.run_in_executor(
                self.executor,
                self._extract_word_content,
                file_path,
                document_id,
                extract_text,
                extract_tables
            )
            return [page_data]
        except Exception as e:
            logger.error(f"Word extraction failed: {str(e)}")
            raise RuntimeError(f"Word extraction failed: {str(e)}")

    def _extract_word_content(
        self,
        file_path: str,
        document_id: str,
        extract_text: bool,
        extract_tables: bool
    ) -> PageData:
        page_data = PageData(page_number=1)

        try:
            doc = Document(file_path)

            if extract_text:
                text_content = ""
                for paragraph in doc.paragraphs:
                    text_content += paragraph.text + "\n"
                page_data.text_content = text_content

            if extract_tables:
                for table_idx, table in enumerate(doc.tables):
                    headers = []
                    rows = []

                    if table.rows:
                        headers = [cell.text for cell in table.rows[0].cells]
                        for row in table.rows[1:]:
                            row_data = [cell.text for cell in row.cells]
                            rows.append(row_data)

                    table_info = TableInfo(
                        table_id=f"{document_id}_word_table_{table_idx}",
                        page_number=1,
                        headers=headers,
                        rows=rows
                    )
                    page_data.tables.append(table_info)

            logger.info(
                f"Word doc extracted: {len(page_data.text_content)} chars, "
                f"{len(page_data.tables)} tables"
            )

        except Exception as e:
            logger.error(f"Word extraction failed: {str(e)}")
            page_data.metadata = {"error": str(e)}

        return page_data


class AsyncCSVExtractor:
    """Async CSV extraction"""

    def __init__(self):
        self.executor = ThreadPoolExecutor(max_workers=settings.MAX_WORKERS)

    async def extract_from_csv(
        self,
        file_bytes: bytes,
        document_id: str
    ) -> List[PageData]:
        """Extract data from CSV"""
        loop = asyncio.get_running_loop()
        try:
            page_data = await loop.run_in_executor(
                self.executor,
                self._parse_csv,
                file_bytes,
                document_id
            )
            return [page_data]
        except Exception as e:
            logger.error(f"CSV extraction failed: {str(e)}")
            raise RuntimeError(f"CSV extraction failed: {str(e)}")

    def _parse_csv(self, file_bytes: bytes, document_id: str) -> PageData:
        page_data = PageData(page_number=1)

        try:
            df = pd.read_csv(BytesIO(file_bytes))

            page_data.text_content = df.to_string()

            headers = [str(h) for h in df.columns.tolist()]
            rows = [[str(cell) for cell in row] for row in df.values.tolist()]

            table_info = TableInfo(
                table_id=f"{document_id}_csv_table",
                page_number=1,
                headers=headers,
                rows=rows
            )
            page_data.tables.append(table_info)

            logger.info(f"CSV extracted: {len(df)} rows, {len(df.columns)} cols")

        except Exception as e:
            logger.error(f"CSV parsing failed: {str(e)}")
            page_data.metadata = {"error": str(e)}

        return page_data


class AsyncTextExtractor:
    """Async plain text extraction"""

    @staticmethod
    async def extract_from_text(file_bytes: bytes) -> List[PageData]:
        """Extract plain text content"""
        page_data = PageData(page_number=1)

        try:
            try:
                text_content = file_bytes.decode('utf-8')
            except UnicodeDecodeError:
                text_content = file_bytes.decode('latin-1', errors='replace')

            page_data.text_content = text_content
            logger.info(f"Text file extracted: {len(text_content)} characters")

        except Exception as e:
            logger.error(f"Text extraction failed: {str(e)}")
            page_data.text_content = f"Error: {str(e)}"

        return [page_data]